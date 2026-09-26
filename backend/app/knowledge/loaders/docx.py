from io import BytesIO

from docx import Document

from app.knowledge.chunking import TextSection
from app.knowledge.errors import PermanentIngestionError


def extract_docx(data: bytes, filename: str) -> list[TextSection]:
    try:
        document = Document(BytesIO(data))
    except PermanentIngestionError:
        raise
    except Exception as exc:
        raise PermanentIngestionError("This document could not be read.") from exc
    parts: list[str] = [
        paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
    ]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise PermanentIngestionError("This document has no extractable text.")
    return [TextSection(text, {"source": filename, "section": "document"})]
